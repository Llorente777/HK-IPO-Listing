"""Import historical workbook structure without treating legacy formulas as current production data."""
from __future__ import annotations
import argparse, json
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("data/imported/excel_history.json"))
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.workbook, data_only=False, read_only=True)
    result = {"source_file": args.workbook.name, "imported_at": datetime.now(timezone.utc).isoformat(), "production_current": False, "sheets": []}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True))
        headers = [str(value).strip() if value is not None else None for value in (rows[0] if rows else [])]
        result["sheets"].append({"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "headers": headers, "note": "Historical structure only; formulas and plugin values are not production evidence."})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

if __name__ == "__main__": main()
